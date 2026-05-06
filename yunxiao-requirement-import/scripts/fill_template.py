from __future__ import annotations

import argparse
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


HEADERS = [
    "标题",
    "描述",
    "状态",
    "负责人",
    "优先级",
    "迭代",
    "参与者",
    "抄送",
    "标签",
    "计划开始时间",
    "计划完成时间",
    "预计工时",
]


def load_items(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    items = data.get("items", data) if isinstance(data, dict) else data
    if not isinstance(items, list):
        raise ValueError("JSON must be a list or an object with an items list")
    return items


def value(item: dict, key: str, default: str = "") -> str:
    val = item.get(key, default)
    return "" if val is None else str(val)


def main() -> None:
    here = Path(__file__).resolve().parent
    default_template = here.parent / "assets" / "产品类需求 导入模板.xlsx"

    parser = argparse.ArgumentParser()
    parser.add_argument("--items", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--template", type=Path, default=default_template)
    parser.add_argument("--owner", default="潘洪森")
    parser.add_argument("--include-description", action="store_true")
    parser.add_argument("--include-tags", action="store_true")
    args = parser.parse_args()

    wb = load_workbook(args.template)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    if headers[: len(HEADERS)] != HEADERS:
        raise ValueError(f"Unexpected template headers: {headers}")

    while ws.max_row > 1:
        ws.delete_rows(2)

    for item in load_items(args.items):
        if not item.get("title"):
            continue
        row = [
            value(item, "title"),
            value(item, "description") if args.include_description else "",
            value(item, "status"),
            value(item, "owner", args.owner),
            value(item, "priority"),
            value(item, "iteration"),
            value(item, "participants"),
            value(item, "cc"),
            value(item, "tags") if args.include_tags else "",
            value(item, "start_date"),
            value(item, "due_date"),
            value(item, "estimate_hours"),
        ]
        ws.append(row)

    base = ws[1]
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row_idx, col_idx)
            ref = base[col_idx - 1]
            cell.font = copy(ref.font) if ref.font else Font(name="Arial")
            if not cell.font.name:
                cell.font = Font(name="Arial")
            cell.fill = copy(ref.fill)
            cell.border = copy(ref.border)
            cell.number_format = copy(ref.number_format)
            cell.protection = copy(ref.protection)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col in ws[1]:
        if col.value == "标题":
            ws.column_dimensions[col.column_letter].width = max(ws.column_dimensions[col.column_letter].width or 0, 42)
        elif col.value == "描述":
            ws.column_dimensions[col.column_letter].width = max(ws.column_dimensions[col.column_letter].width or 0, 28)
        else:
            ws.column_dimensions[col.column_letter].width = max(ws.column_dimensions[col.column_letter].width or 0, 13)

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 34

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
